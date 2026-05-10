"""CV generator — full or job-tailored, exports PDF/DOCX/Markdown."""
from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Any, TYPE_CHECKING

from src.config import settings
from src.skills.base import BaseSkill, SkillResponse

if TYPE_CHECKING:
    from src.brain.context_builder import Context
    from src.brain.llm_client import LLMClient
    from src.memory.manager import MemoryManager


logger = logging.getLogger(__name__)


CV_PROMPT = """You are an expert CV writer. Build a professional CV in Markdown.

User profile:
{profile}

Work experience:
{work}

Education:
{education}

Projects:
{projects}

Skills:
{skills}

Achievements:
{achievements}

{job_description_block}

Output a clean Markdown CV with sections:
# {name}
**Contact line** (omit fields you don't have)

## Summary
2-3 sentences. {summary_focus}

## Experience
For each role: **Role — Company** | dates
- 2-4 impact bullets focused on outcomes and metrics.

## Education
For each: **Degree, Field — Institution** | dates

## Projects
For each: **Project Name** — short summary, key tech, impact.

## Skills
Group by category.

## Achievements
Bulleted list.

Be specific. Don't invent. {tailoring_instruction}
"""


class CVGeneratorSkill(BaseSkill):
    name = "cv_generator"
    description = "Build a CV — either full or tailored to a specific job description."
    triggers = [
        "/cv", "/cv_for", "cv", "curriculum", "resume", "résumé",
        "generate my cv", "tailored cv", "make a cv", "build a cv",
        "hazme un cv", "genera mi cv", "adapta mi cv",
    ]
    examples = [
        "generate my CV tailored for this job: [paste job description]",
        "hazme un cv para una posición de data scientist en Glovo",
    ]
    input_schema = {
        "type": "object",
        "properties": {
            "job_description": {
                "type": "string",
                "description": "Optional job description to tailor the CV to. If empty, generates a full general CV.",
            },
        },
        "required": [],
    }

    def __init__(self, llm: "LLMClient", memory: "MemoryManager") -> None:
        self.llm = llm
        self.memory = memory

    def can_handle(self, message: str, intent: str | None = None) -> float:
        m = message.lower()
        if m.startswith("/cv"):
            return 0.95
        if any(t in m for t in ("curriculum", "resume", "résumé")):
            return 0.7
        if "cv" in m and any(k in m for k in ("generate", "make", "build", "tailor", "write")):
            return 0.8
        return super().can_handle(message, intent)

    async def execute(self, message: str, context: "Context") -> SkillResponse:
        job_description = self._extract_job_description(message)
        profile = context.profile or {}

        markdown = await self._render_cv(profile=profile, job_description=job_description)

        export_dir = settings.exports_dir
        export_dir.mkdir(parents=True, exist_ok=True)
        ts = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
        slug = self._slug(profile.get("name") or "cv")
        suffix = "_tailored" if job_description else ""

        md_path = export_dir / f"{slug}{suffix}_{ts}.md"
        md_path.write_text(markdown, encoding="utf-8")

        files: list[dict[str, Any]] = [{"path": str(md_path), "type": "markdown", "name": md_path.name}]

        try:
            docx_path = export_dir / f"{slug}{suffix}_{ts}.docx"
            self._write_docx(markdown, docx_path)
            files.append({"path": str(docx_path), "type": "docx", "name": docx_path.name})
        except Exception as exc:
            logger.warning("DOCX export failed: %s", exc)

        try:
            pdf_path = export_dir / f"{slug}{suffix}_{ts}.pdf"
            self._write_pdf(markdown, pdf_path)
            files.append({"path": str(pdf_path), "type": "pdf", "name": pdf_path.name})
        except Exception as exc:
            logger.warning("PDF export failed: %s", exc)

        text = (
            f"CV ready{' (tailored)' if job_description else ''}. "
            f"Generated {len(files)} format(s)."
        )
        return SkillResponse(
            text=text,
            files=files,
            metadata={"tailored": bool(job_description)},
        )

    @staticmethod
    def _extract_job_description(message: str) -> str | None:
        m = message.strip()
        if m.lower().startswith("/cv_for"):
            jd = m.split(maxsplit=1)
            return jd[1].strip() if len(jd) > 1 else None
        # Heuristic: long-ish message with job-y keywords
        if len(m) > 200 and any(k in m.lower() for k in (
            "responsibilities", "requirements", "we're looking", "we are looking",
            "job posting", "qualifications",
        )):
            return m
        return None

    async def _render_cv(self, profile: dict, job_description: str | None) -> str:
        prompt = CV_PROMPT.format(
            profile=self._fmt_profile(profile),
            work=self._fmt_list(profile.get("work_experience"), self._fmt_job),
            education=self._fmt_list(profile.get("education"), self._fmt_education),
            projects=self._fmt_list(profile.get("projects"), self._fmt_project),
            skills=", ".join(s["name"] for s in (profile.get("skills") or []))[:1000] or "(none)",
            achievements="(use stored achievements if known)",
            job_description_block=(
                f"\nTarget role / job description:\n\"\"\"\n{job_description}\n\"\"\"\n"
                if job_description else ""
            ),
            name=profile.get("name") or "Your Name",
            summary_focus=(
                "Tailor it to the target role." if job_description
                else "Lead with the user's strongest evidence."
            ),
            tailoring_instruction=(
                "Reorder and emphasise items that match the target job. "
                "Drop or shorten unrelated content."
                if job_description else "Keep it balanced and comprehensive."
            ),
        )
        return await self.llm.complete(
            prompt=prompt,
            system="You are a senior career writer. Output Markdown only.",
            model_tier="powerful",
            max_tokens=3000,
        )

    @staticmethod
    def _fmt_profile(p: dict) -> str:
        return (
            f"Name: {p.get('name')}\n"
            f"Bio: {p.get('bio') or p.get('living_profile') or '(unknown)'}\n"
            f"Communication style: {p.get('communication_style') or '(neutral)'}"
        )

    @staticmethod
    def _fmt_list(items: list | None, fmt) -> str:
        if not items:
            return "(none)"
        return "\n".join(f"- {fmt(i)}" for i in items)

    @staticmethod
    def _fmt_job(j: dict) -> str:
        return f"{j.get('role')} @ {j.get('company')} ({j.get('start_date')} → {j.get('end_date') or 'present'})"

    @staticmethod
    def _fmt_education(e: dict) -> str:
        return f"{e.get('degree') or ''} {e.get('field') or ''} @ {e.get('institution')}"

    @staticmethod
    def _fmt_project(p: dict) -> str:
        return (
            f"{p.get('name')} ({p.get('status')}) — {p.get('description') or ''}"
        )

    @staticmethod
    def _slug(s: str) -> str:
        return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_") or "cv"

    # --- Exporters ----------------------------------------------------
    @staticmethod
    def _write_docx(markdown: str, path: Path) -> None:
        from docx import Document

        doc = Document()
        for line in markdown.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith(("- ", "* ")):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(path)

    @staticmethod
    def _write_pdf(markdown: str, path: Path) -> None:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=50, rightMargin=50)
        story = []
        for line in markdown.splitlines():
            if not line.strip():
                story.append(Spacer(1, 6))
                continue
            if line.startswith("# "):
                story.append(Paragraph(line[2:], styles["Title"]))
            elif line.startswith("## "):
                story.append(Paragraph(line[3:], styles["Heading2"]))
            elif line.startswith("### "):
                story.append(Paragraph(line[4:], styles["Heading3"]))
            elif line.startswith(("- ", "* ")):
                story.append(Paragraph("• " + line[2:], styles["BodyText"]))
            else:
                story.append(Paragraph(line, styles["BodyText"]))
        doc.build(story)
